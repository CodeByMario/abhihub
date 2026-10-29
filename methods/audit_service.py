"""
Admin Audit & Transaction Export Service
Generates structured multi-sheet Excel (.xlsx) workbooks for admin financial audits.
"""
import io
import logging
from datetime import datetime, timezone
from typing import Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from methods.supabase_helper import init_supabase, supabase

log = logging.getLogger('abhihub.audit_service')


def _db():
    from methods.supabase_helper import init_supabase_admin
    return init_supabase_admin() or init_supabase() or supabase


def _style_header(ws, col_count: int):
    """Applies styled headers with dark background and white bold text."""
    fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
    ws.row_dimensions[1].height = 26


def _auto_fit_columns(ws):
    """Auto-adjusts column widths to fit content cleanly."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)


def build_audit_export(filters: Optional[dict] = None) -> io.BytesIO:
    """
    Builds a multi-sheet audit Excel workbook based on optional filters:
      - start_date (ISO string or YYYY-MM-DD)
      - end_date (ISO string or YYYY-MM-DD)
      - user_id (optional creator/user filter)
      - status (optional status filter)
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl is not installed")

    filters = filters or {}
    start_date = filters.get('start_date')
    end_date = filters.get('end_date')
    user_id = filters.get('user_id')
    status = filters.get('status')

    wb = openpyxl.Workbook()
    # Remove default empty sheet
    wb.remove(wb.active)

    db = _db()

    # ── Sheet 1: Razorpay Orders / Credit Transactions ─────────────────────────
    ws1 = wb.create_sheet(title="Razorpay Orders")
    headers1 = ["ID", "User ID", "Order ID", "Amount (INR)", "UTR", "Status", "Client IP", "Admin Notes", "Created At"]
    ws1.append(headers1)

    if db:
        try:
            q = db.table('credit_transactions').select('*')
            if user_id:
                q = q.eq('user_id', user_id)
            if status:
                q = q.eq('status', status)
            if start_date:
                q = q.gte('created_at', start_date)
            if end_date:
                q = q.lte('created_at', end_date)
            res1 = q.order('created_at', desc=True).limit(2000).execute()
            for r in (res1.data or []):
                ws1.append([
                    str(r.get('id', '')),
                    str(r.get('user_id', '')),
                    str(r.get('order_id', '')),
                    float(r.get('amount', 0)),
                    str(r.get('utr', '') or ''),
                    str(r.get('status', '')),
                    str(r.get('client_ip', '') or ''),
                    str(r.get('admin_notes', '') or ''),
                    str(r.get('created_at', ''))
                ])
        except Exception as e:
            log.warning(f"[AUDIT] Sheet 1 fetch error: {e}")

    _style_header(ws1, len(headers1))
    _auto_fit_columns(ws1)

    # ── Sheet 2: Creator Payout Requests ───────────────────────────────────────
    ws2 = wb.create_sheet(title="Creator Payouts")
    headers2 = ["Payout ID", "Creator ID", "Amount (INR)", "Method", "Status", "Processed At", "Admin Notes", "Created At"]
    ws2.append(headers2)

    if db:
        try:
            q = db.table('creator_payout_requests').select('*')
            if user_id:
                q = q.eq('creator_id', user_id)
            if status:
                q = q.eq('status', status)
            if start_date:
                q = q.gte('created_at', start_date)
            if end_date:
                q = q.lte('created_at', end_date)
            res2 = q.order('created_at', desc=True).limit(2000).execute()
            for r in (res2.data or []):
                ws2.append([
                    str(r.get('id', '')),
                    str(r.get('creator_id', '')),
                    float(r.get('amount', 0)),
                    str(r.get('method', '') or ''),
                    str(r.get('status', '')),
                    str(r.get('processed_at', '') or ''),
                    str(r.get('admin_notes', '') or ''),
                    str(r.get('created_at', ''))
                ])
        except Exception as e:
            log.warning(f"[AUDIT] Sheet 2 fetch error: {e}")

    _style_header(ws2, len(headers2))
    _auto_fit_columns(ws2)

    # ── Sheet 3: Creator Wallet Ledger (INR) ───────────────────────────────────
    ws3 = wb.create_sheet(title="Wallet Ledger")
    headers3 = ["ID", "Creator ID", "Amount (INR)", "Type", "Status", "Order ID", "Unlock At", "Created At"]
    ws3.append(headers3)

    if db:
        try:
            q = db.table('wallet_ledger').select('*')
            if user_id:
                q = q.eq('creator_id', user_id)
            if status:
                q = q.eq('status', status)
            if start_date:
                q = q.gte('created_at', start_date)
            if end_date:
                q = q.lte('created_at', end_date)
            res3 = q.order('created_at', desc=True).limit(2000).execute()
            for r in (res3.data or []):
                ws3.append([
                    str(r.get('id', '')),
                    str(r.get('creator_id', '')),
                    float(r.get('amount', 0)),
                    str(r.get('type', '')),
                    str(r.get('status', '')),
                    str(r.get('order_id', '') or ''),
                    str(r.get('unlock_at', '') or ''),
                    str(r.get('created_at', ''))
                ])
        except Exception as e:
            log.warning(f"[AUDIT] Sheet 3 fetch error: {e}")

    _style_header(ws3, len(headers3))
    _auto_fit_columns(ws3)

    # ── Sheet 4: Unified View Credit Ledger ────────────────────────────────────
    ws4 = wb.create_sheet(title="View Credit Ledger")
    headers4 = ["ID", "User ID", "Delta", "Balance", "Reason", "Source", "Reference ID", "Created At"]
    ws4.append(headers4)

    if db:
        try:
            q = db.table('user_credit_ledger').select('*')
            if user_id:
                q = q.eq('user_id', user_id)
            if start_date:
                q = q.gte('created_at', start_date)
            if end_date:
                q = q.lte('created_at', end_date)
            res4 = q.order('created_at', desc=True).limit(2000).execute()
            for r in (res4.data or []):
                ws4.append([
                    str(r.get('id', '')),
                    str(r.get('user_id', '')),
                    int(r.get('delta', 0)),
                    int(r.get('balance', 0)),
                    str(r.get('reason', '')),
                    str(r.get('source', '') or ''),
                    str(r.get('ref_id', '') or ''),
                    str(r.get('created_at', ''))
                ])
        except Exception as e:
            log.warning(f"[AUDIT] Sheet 4 fetch error: {e}")

    _style_header(ws4, len(headers4))
    _auto_fit_columns(ws4)

    # ── Sheet 5: KYC Records ───────────────────────────────────────────────────
    ws5 = wb.create_sheet(title="KYC Records")
    headers5 = ["ID", "User ID", "Full Name", "PAN (Masked)", "Method", "UPI ID", "Bank Name", "IFSC", "Verified", "Submitted At"]
    ws5.append(headers5)

    if db:
        try:
            res5 = db.table('user_kyc_details').select('*').order('submitted_at', desc=True).limit(1000).execute()
            for r in (res5.data or []):
                pan = str(r.get('pan_number') or '')
                masked_pan = (pan[:5] + '****' + pan[-1]) if len(pan) == 10 else (pan[:2] + '****') if pan else ''
                ws5.append([
                    str(r.get('id', '')),
                    str(r.get('user_id', '')),
                    str(r.get('full_name', '')),
                    masked_pan,
                    str(r.get('preferred_method', '')),
                    str(r.get('upi_id', '') or ''),
                    str(r.get('bank_name', '') or ''),
                    str(r.get('ifsc_code', '') or ''),
                    "YES" if r.get('is_verified') else "NO",
                    str(r.get('submitted_at', ''))
                ])
        except Exception as e:
            log.warning(f"[AUDIT] Sheet 5 fetch error: {e}")

    _style_header(ws5, len(headers5))
    _auto_fit_columns(ws5)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
